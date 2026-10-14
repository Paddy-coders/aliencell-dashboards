# -*- coding: utf-8 -*-
"""
将用户正式目标表 Excel 转换为 monthly_targets.json
- 纯收入(人民币) → sales (目标表要求人民币)
- 广告费%(小数)  → adShare (百分比数值, 如 0.06 → 6)
- 月销量         → units
- 一级站点 US/EU/UK → country (target 页 regionOf 自动映射 US→NA)
- 全零行(销量/纯收入/广告全 0)跳过
"""
import openpyxl, json

SRC = r"C:\Users\paddy\Desktop\Aliencell 海外电商2026年度目标（E1低价版）.xlsx"
OUT = r"C:\Users\paddy\WorkBuddy\2026-08-18-14-43-38\monthly_targets.json"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb['亚马逊']

targets = []
skipped_zero = 0
for r in range(2, ws.max_row + 1):
    code = ws.cell(row=r, column=1).value
    if not code or not str(code).strip():
        continue
    site = (str(ws.cell(row=r, column=2).value) or '').strip()
    product = (str(ws.cell(row=r, column=3).value) or '').strip()
    if not site or not product:
        continue
    for m in range(1, 13):
        units = ws.cell(row=r, column=4 + m).value or 0        # 月销量
        sales = ws.cell(row=r, column=53 + m).value or 0        # 月纯收入(人民币)
        adpct = ws.cell(row=r, column=41 + m).value             # 月广告费%(小数)
        try: units = float(units)
        except (TypeError, ValueError): units = 0
        try: sales = float(sales)
        except (TypeError, ValueError): sales = 0
        try: adpct = float(adpct)
        except (TypeError, ValueError): adpct = 0
        if units == 0 and sales == 0 and adpct == 0:
            skipped_zero += 1
            continue
        targets.append({
            'month': '2026-%02d' % m,
            'country': site,
            'product': product,
            'units': units if units == int(units) else round(units, 1),
            'sales': round(sales, 2),
            'adShare': round(adpct * 100, 2)
        })

out = {
    'sample': False,
    'note': '正式目标表: Aliencell 海外电商2026年度目标（E1低价版）.xlsx 亚马逊Sheet; 纯收入(人民币)→销售额, 广告费%→广告占比, 月销量→销量; 全零行已跳过',
    'targets': targets
}
with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)

# 统计
from collections import Counter
months = Counter(t['month'] for t in targets)
sites = Counter(t['country'] for t in targets)
print('目标行数:', len(targets), '(跳过全零:', skipped_zero, ')')
print('按月:', dict(sorted(months.items())))
print('按站点:', dict(sites))
print('产品数:', len({(t['country'], t['product']) for t in targets}))
# 8月样例
print('\n8月样例(前5条):')
for t in [t for t in targets if t['month'] == '2026-08'][:5]:
    print(' ', t)
print('\n12月样例(前5条):')
for t in [t for t in targets if t['month'] == '2026-12'][:5]:
    print(' ', t)
